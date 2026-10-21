from datetime import datetime, timezone
from io import BytesIO
import threading
import zipfile

from docx import Document

from app.ai.llm import LLMError
from app.services.JobSearch import JOBS_PER_AREA, analyze_cv_stream
from conftest import analysis_from_response


SOFTWARE_PROFILE = {
    "candidate_name": None,
    "areas": [
        {
            "key": "software_engineering",
            "name": "Engenharia de Software",
            "search_query": "software engineer python react",
            "search_tag": "software",
            "industries": ["software engineering"],
            "role_terms": ["software engineer", "engenheira de software", "software developer", "programador", "desenvolvedor"],
            "matched_skills": ["python", "react"],
            "confidence": "high",
            "evidence": [
                {"text": "Engenheira de Software com Python e React", "section": "experiencia"},
                {"text": "Licenciatura em Engenharia Informática", "section": "formacao"},
            ],
        }
    ],
}


def use_software_profile(monkeypatch):
    monkeypatch.setattr("app.services.CareerProfile.chat_json", lambda *args, **kwargs: SOFTWARE_PROFILE)


def stub_job_batches(monkeypatch, jobs, search_calls=None):
    def fake(tag, location=None, limit=10, include_remote=True):
        if search_calls is not None:
            search_calls.append({
                "tag": tag, "location": location, "limit": limit, "include_remote": include_remote,
            })
        yield "Jobicy", jobs() if callable(jobs) else jobs

    monkeypatch.setattr("app.services.JobSearch.fetch_source_batches", fake)


def create_cv_docx() -> bytes:
    document = Document()
    document.add_paragraph("Experiência Profissional")
    document.add_paragraph("Engenheira de Software com Python e React")
    document.add_paragraph("Formação Académica")
    document.add_paragraph("Licenciatura em Engenharia Informática")
    content = BytesIO()
    document.save(content)
    return content.getvalue()


def test_cv_upload_detects_area_and_searches_with_cv_terms(client, monkeypatch):
    search_calls = []

    def fake_jobs():
        return [
            {
                "company": "Critical Manufacturing",
                "title": "Software Engineer",
                "url": "https://jobicy.com/jobs/example",
                "location": "Portugal",
                "industry": ["Software Engineering"],
                "excerpt": "Python and React experience",
                "source": "Jobicy",
            },
            {
                "company": "Example US company",
                "title": "Senior Software Engineer",
                "url": "https://jobicy.com/jobs/us-example",
                "location": "Portugal, USA",
                "industry": ["Software Engineering"],
                "excerpt": "Python and React experience",
                "source": "Jobicy",
            },
            {
                "company": "QuickNode",
                "title": "Software Engineer",
                "url": "https://jobicy.com/jobs/quicknode-portugal",
                "location": "Portugal",
                "industry": ["Software Engineering"],
                "excerpt": "Python and React experience",
                "source": "Jobicy",
            },
            {
                "company": "QuickNode",
                "title": "Senior Software Engineer",
                "url": "https://jobicy.com/jobs/quicknode-portugal-usa",
                "location": "Portugal, USA",
                "industry": ["Software Engineering"],
                "excerpt": "Python and React experience",
                "source": "Jobicy",
            },
            {
                "company": "Europe company",
                "title": "Senior Software Engineer",
                "url": "https://jobicy.com/jobs/europe-example",
                "location": "Europe",
                "industry": ["Software Engineering"],
                "excerpt": "Python and React experience",
                "source": "Jobicy",
            },
            {
                "company": "Global company",
                "title": "Senior Software Engineer",
                "url": "https://jobicy.com/jobs/global-example",
                "location": "Remote, Worldwide",
                "industry": ["Software Engineering"],
                "excerpt": "Python and React experience",
                "source": "Jobicy",
            },
            {
                "company": "US remote company",
                "title": "Senior Software Engineer",
                "url": "https://jobicy.com/jobs/remote-us-example",
                "location": "Remote, USA",
                "industry": ["Software Engineering"],
                "excerpt": "Python and React experience",
                "source": "Jobicy",
            },
            {
                "company": "US-only company",
                "title": "Senior Software Engineer",
                "url": "https://jobicy.com/jobs/us-only-example",
                "location": "United States",
                "industry": ["Software Engineering"],
                "excerpt": "Python and React experience",
                "source": "Jobicy",
            },
        ]

    use_software_profile(monkeypatch)
    stub_job_batches(monkeypatch, fake_jobs, search_calls)

    response = client.post(
        "/career-profile/analyze",
        data={"location": "Portugal"},
        files={
            "file": (
                "cv.docx",
                create_cv_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    result = analysis_from_response(response)
    assert result["cv_saved"] is True
    assert result["areas"][0]["name"] == "Engenharia de Software"
    assert {evidence["section"] for evidence in result["areas"][0]["evidence"]} == {
        "experiencia",
        "formacao",
    }
    assert "python" in result["areas"][0]["search_query"].casefold()
    assert "text/event-stream" in response.headers.get("content-type", "")
    assert search_calls[0]["tag"] == "engenharia de software"
    assert search_calls[0]["location"] == "Portugal"
    assert search_calls[0]["limit"] == JOBS_PER_AREA
    assert result["llm_used"] is True
    assert [suggestion["company"] for suggestion in result["suggestions"]] == [
        "Critical Manufacturing",
        "QuickNode",
        "Example US company",
        "Europe company",
        "Global company",
    ]
    assert [suggestion["location_priority"] for suggestion in result["suggestions"]] == [
        0,
        0,
        1,
        2,
        3,
    ]
    assert result["suggestions"][2]["location"] == "Portugal, USA"
    assert "US remote company" not in [suggestion["company"] for suggestion in result["suggestions"]]
    assert "US-only company" not in [suggestion["company"] for suggestion in result["suggestions"]]
    assert result["excluded_location_count"] == 2
    assert result["suggestions"][1]["url"] == "https://jobicy.com/jobs/quicknode-portugal"
    assert result["suggestions"][0]["email"] is None
    assert len(result["location_priority_order"]) == 5
    assert result["suggestions"][0]["source"] == "Jobicy"


def test_stored_profile_skips_llm(monkeypatch):
    calls = []
    monkeypatch.setattr(
        "app.services.CareerProfile.chat_json",
        lambda *args, **kwargs: calls.append(1) or SOFTWARE_PROFILE,
    )
    stub_job_batches(monkeypatch, [])
    events = list(analyze_cv_stream(
        "cv.docx",
        create_cv_docx(),
        profile={**SOFTWARE_PROFILE, "languages": ["Português"], "llm_used": True},
    ))
    assert calls == []
    assert events[0]["type"] == "profile"
    assert events[0]["llm_used"] is True
    assert events[0]["areas"][0]["name"] == "Engenharia de Software"
    assert events[0]["languages"] == ["Português"]


def test_analyze_stored_reuses_saved_profile(client, monkeypatch):
    calls = []

    def fake_llm(*args, **kwargs):
        calls.append(1)
        return SOFTWARE_PROFILE

    monkeypatch.setattr("app.services.CareerProfile.chat_json", fake_llm)
    stub_job_batches(monkeypatch, [])
    first = client.post(
        "/career-profile/analyze",
        data={"location": "Portugal"},
        files={
            "file": (
                "cv.docx",
                create_cv_docx(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert first.status_code == 200
    assert calls == [1]
    again = client.post("/career-profile/analyze-stored", json={"location": "Lisboa"})
    assert again.status_code == 200
    assert analysis_from_response(again)["areas"][0]["name"] == "Engenharia de Software"
    assert calls == [1]


def test_edited_profile_drives_the_next_search(client, monkeypatch):
    use_software_profile(monkeypatch)
    stub_job_batches(monkeypatch, [])
    client.post(
        "/career-profile/analyze",
        data={"location": "Portugal"},
        files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    edited = client.put("/career-profile/profile", json={"areas": [
        {"name": "Engenharia de Software", "skills": ["Python", "Django"]},
        {"name": "Data Engineer", "skills": ["SQL", "Airflow"]},
    ]})
    assert edited.status_code == 200
    kept, renamed = edited.json()["areas"]
    assert kept["matched_skills"] == ["Python", "Django"]
    assert "software developer" in kept["role_terms"]
    assert renamed["role_terms"] == ["Data Engineer"]
    assert renamed["industries"] == []
    stored = client.get("/career-profile/cv").json()
    assert stored["skills"] == ["Python", "Django", "SQL", "Airflow"]
    assert [area["name"] for area in stored["areas"]] == ["Engenharia de Software", "Data Engineer"]

    searched = []
    stub_job_batches(monkeypatch, [], searched)
    result = analysis_from_response(client.post("/career-profile/analyze-stored", json={"location": "Portugal"}))
    assert [area["name"] for area in result["areas"]] == ["Engenharia de Software", "Data Engineer"]
    assert result["llm_used"] is True
    assert sorted(call["tag"] for call in searched) == ["data engineer", "engenharia de software"]


def test_profile_edit_validates_areas(client):
    assert client.put("/career-profile/profile", json={"areas": [{"name": "Dev", "skills": []}]}).status_code == 404
    assert client.put("/career-profile/profile", json={"areas": []}).status_code == 422
    too_many = [{"name": f"Área {index}", "skills": []} for index in range(4)]
    assert client.put("/career-profile/profile", json={"areas": too_many}).status_code == 422


def test_analyze_stream_yields_jobs_before_later_sources(monkeypatch):
    use_software_profile(monkeypatch)
    release_slow = threading.Event()

    def fake(tag, location=None, limit=10, include_remote=True):
        yield "Jobicy", [{
            "company": "Fast Co",
            "title": "Software Engineer",
            "url": "https://x/fast",
            "location": "Portugal",
            "industry": ["Software Engineering"],
            "excerpt": "Python and React",
            "source": "Jobicy",
        }]
        release_slow.wait(timeout=2)
        yield "Slow", [{
            "company": "Slow Co",
            "title": "Software Engineer",
            "url": "https://x/slow",
            "location": "Portugal",
            "industry": ["Software Engineering"],
            "excerpt": "Python and React",
            "source": "Slow",
        }]

    monkeypatch.setattr("app.services.JobSearch.fetch_source_batches", fake)
    events = analyze_cv_stream("cv.docx", create_cv_docx())
    assert next(events)["type"] == "profile"
    first_jobs = next(events)
    assert first_jobs["type"] == "jobs"
    assert [item["company"] for item in first_jobs["suggestions"]] == ["Fast Co"]
    release_slow.set()
    rest = list(events)
    assert any(
        event["type"] == "jobs" and any(item["company"] == "Slow Co" for item in event["suggestions"])
        for event in rest
    )
    assert rest[-1]["type"] == "done"


def test_cv_upload_rejects_unsupported_file_type(client):
    response = client.post(
        "/career-profile/analyze",
        files={"file": ("cv.xlsx", b"PK\x03\x04not-a-cv", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )

    assert response.status_code == 400
    assert "não suportado" in response.json()["detail"].casefold()


def test_cv_upload_accepts_plain_text(client, monkeypatch):
    use_software_profile(monkeypatch)
    stub_job_batches(monkeypatch, [])
    result = analysis_from_response(client.post(
        "/career-profile/analyze",
        data={"location": "Portugal"},
        files={"file": ("cv.txt", "Engenheira de Software com Python e React\n".encode(), "text/plain")},
    ))
    assert result["areas"]


def test_extract_cv_text_reads_odt_and_rtf():
    from app.services.CareerProfile import extract_cv_text

    xml = (
        '<?xml version="1.0"?><office:document-content '
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
        "<office:body><office:text>"
        "<text:p>Engenheira de Software com Python</text:p>"
        "</office:text></office:body></office:document-content>"
    )
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
        archive.writestr("content.xml", xml)

    assert "Python" in extract_cv_text("cv.odt", buffer.getvalue())
    assert "Python" in extract_cv_text("cv.rtf", br"{\rtf1 Engenheira de Software com Python}")


def test_extract_cv_text_reads_jpg(monkeypatch):
    from app.services import CareerProfile as career

    class FakeOcr:
        def __call__(self, content):
            assert content.startswith(b"\xff\xd8\xff")
            return type("Result", (), {"txts": ("Miguel Silva", "Carpinteiro com 10 anos")})()

    monkeypatch.setattr(career, "_ocr_engine", lambda: FakeOcr())
    text = career.extract_cv_text("cv.jpg", b"\xff\xd8\xff\xe0fake-jpeg")
    assert "Carpinteiro" in text
    assert "Miguel Silva" in text


def test_cv_keeps_area_jobs_even_without_skills_in_the_ad(client, monkeypatch):
    use_software_profile(monkeypatch)
    stub_job_batches(monkeypatch, [
        {"company": "Match", "title": "Backend Software Engineer", "url": "https://x/1", "location": "Portugal",
         "industry": [], "excerpt": "We use Python daily", "source": "Jobicy"},
        {"company": "Local", "title": "Programador", "url": "https://x/2", "location": "Lisboa",
         "industry": [], "excerpt": "Empresa: Local<br />Localidade: Lisboa", "source": "ITJobs"},
        {"company": "NoSkills", "title": "Software Engineer", "url": "https://x/3", "location": "Portugal",
         "industry": [], "excerpt": "Cobol mainframes", "source": "Jobicy"},
        {"company": "Wood", "title": "Carpinteiro", "url": "https://x/4", "location": "Porto",
         "industry": [], "excerpt": "Trabalho em madeira", "source": "ITJobs"},
    ])
    response = client.post(
        "/career-profile/analyze",
        data={"location": "Portugal"},
        files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    result = analysis_from_response(response)
    companies = [suggestion["company"] for suggestion in result["suggestions"]]
    assert companies == ["Match", "NoSkills", "Local"]
    assert "python" in result["suggestions"][0]["matched_skills"]
    assert "Wood" not in companies


def test_spoken_requirements_reads_title_not_a_language_list():
    from app.services.JobSearch import _spoken_requirements

    assert _spoken_requirements(".Net Developer (French Speaker)") == ["french"]
    assert _spoken_requirements("Fullstack Developer (f/m/d) Mid") == []
    assert _spoken_requirements("German-speaking Support Engineer") == ["german"]
    assert _spoken_requirements("Comercial falante de frances") == ["frances"]
    assert _spoken_requirements("Native speaker") == []


def test_cv_skips_jobs_that_require_a_language_absent_from_the_cv(client, monkeypatch):
    use_software_profile(monkeypatch)
    stub_job_batches(monkeypatch, [
        {"company": "We are META", "title": ".Net Developer (French Speaker)", "url": "https://x/fr",
         "location": "Porto", "industry": ["software engineering"], "excerpt": "Python and React", "source": "Net-Empregos"},
        {"company": "Local", "title": "Software Engineer", "url": "https://x/ok",
         "location": "Lisboa", "industry": ["software engineering"], "excerpt": "Python", "source": "ITJobs"},
    ])
    result = analysis_from_response(client.post(
        "/career-profile/analyze",
        data={"location": "Portugal"},
        files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ))
    companies = [suggestion["company"] for suggestion in result["suggestions"]]
    assert companies == ["Local"]
    assert "We are META" not in companies


def test_braga_search_ranks_onsite_ahead_of_remote(client, monkeypatch):
    use_software_profile(monkeypatch)
    stub_job_batches(monkeypatch, [
        {"company": "Jobicy Remote", "title": "Senior Software Engineer", "url": "https://jobicy.com/1",
         "location": "Anywhere in the World", "work_mode": "remote", "industry": ["software"],
         "excerpt": "Remote OK", "source": "Jobicy"},
        {"company": "Integer Consulting", "title": "Senior Frontend Developer", "url": "https://www.itjobs.pt/oferta/1",
         "location": "Braga", "industry": ["Full-time"], "excerpt": "Full-time", "source": "ITJobs"},
        {"company": "Lisboa Co", "title": "Software Engineer", "url": "https://www.itjobs.pt/oferta/2",
         "location": "Lisboa", "industry": ["Full-time"], "excerpt": "Full-time", "source": "ITJobs"},
    ])
    result = analysis_from_response(client.post(
        "/career-profile/analyze",
        data={"location": "Braga, Portugal"},
        files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ))
    companies = [suggestion["company"] for suggestion in result["suggestions"]]
    assert companies[0] == "Integer Consulting"
    assert "Jobicy Remote" in companies
    assert "Lisboa Co" not in companies

    without_remote = analysis_from_response(client.post(
        "/career-profile/analyze",
        data={"location": "Braga, Portugal", "work_mode": "onsite"},
        files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ))
    assert [suggestion["company"] for suggestion in without_remote["suggestions"]] == ["Integer Consulting"]


def test_same_company_from_two_boards_keeps_one_card(client, monkeypatch):
    use_software_profile(monkeypatch)
    stub_job_batches(monkeypatch, [
        {"company": "Growin - Know to grow", "title": "Fullstack Developer", "url": "https://www.itjobs.pt/oferta/1",
         "location": "Braga", "industry": ["Full-time"], "excerpt": "Full-time", "source": "ITJobs"},
        {"company": "Growin", "title": "Fullstack Developer", "url": "https://pt.teamlyzer.com/companies/growin/job/1",
         "location": "Braga (Regime: Híbrido)", "industry": ["Backend"], "excerpt": "Backend", "source": "Teamlyzer"},
        {"company": "Checkmarx Portugal", "title": "Senior software engineer", "url": "https://pt.teamlyzer.com/companies/checkmarx/job/1",
         "location": "Braga (Regime: Híbrido)", "industry": ["Backend"], "excerpt": "Backend", "source": "Teamlyzer"},
    ])
    result = analysis_from_response(client.post(
        "/career-profile/analyze",
        data={"location": "Braga, Portugal"},
        files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ))
    companies = [suggestion["company"] for suggestion in result["suggestions"]]
    assert "Checkmarx Portugal" in companies
    assert sum(1 for name in companies if "growin" in name.casefold()) == 1


def test_portugal_keeps_the_same_company_in_two_cities(client, monkeypatch):
    use_software_profile(monkeypatch)
    stub_job_batches(monkeypatch, [
        {"company": "Bosch", "title": "Software Engineer", "url": "https://www.itjobs.pt/oferta/lisboa",
         "location": "Lisboa", "industry": ["Full-time"], "excerpt": "Python", "source": "ITJobs"},
        {"company": "Bosch", "title": "Software Engineer", "url": "https://www.itjobs.pt/oferta/braga",
         "location": "Braga", "industry": ["Full-time"], "excerpt": "Python", "source": "ITJobs"},
    ])
    result = analysis_from_response(client.post(
        "/career-profile/analyze",
        data={"location": "Portugal"},
        files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ))
    places = [suggestion["location"] for suggestion in result["suggestions"] if suggestion["company"] == "Bosch"]
    assert sorted(places) == ["Braga", "Lisboa"]


def test_onsite_search_skips_remote_boards(client, monkeypatch):
    search_calls = []
    use_software_profile(monkeypatch)
    stub_job_batches(monkeypatch, [
        {"company": "Local", "title": "Software Engineer", "url": "https://www.itjobs.pt/oferta/1",
         "location": "Lisboa", "industry": ["software engineering"], "excerpt": "Python", "source": "ITJobs"},
    ], search_calls)
    analysis_from_response(client.post(
        "/career-profile/analyze",
        data={"location": "Portugal", "work_mode": "onsite"},
        files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ))
    assert search_calls[0]["include_remote"] is False


def test_part_time_filter_drops_full_time_and_unknown(client, monkeypatch):
    use_software_profile(monkeypatch)
    stub_job_batches(monkeypatch, [
        {"company": "Inteiro", "title": "Software Engineer", "url": "https://x/full",
         "location": "Lisboa", "industry": ["software engineering"], "excerpt": "Python",
         "source": "ITJobs", "schedule": "full_time"},
        {"company": "Parcial", "title": "Software Engineer part-time", "url": "https://x/part",
         "location": "Lisboa", "industry": ["software engineering"], "excerpt": "Python",
         "source": "ITJobs", "schedule": "part_time"},
        {"company": "Sem horário", "title": "Software Engineer", "url": "https://x/none",
         "location": "Lisboa", "industry": ["software engineering"], "excerpt": "Python",
         "source": "ITJobs"},
    ])
    result = analysis_from_response(client.post(
        "/career-profile/analyze",
        data={"location": "Portugal", "schedule": "part_time"},
        files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ))
    assert [suggestion["company"] for suggestion in result["suggestions"]] == ["Parcial"]
    assert result["suggestions"][0]["schedule"] == "part_time"


def test_suggestion_carries_date_work_mode_and_short_summary(client, monkeypatch):
    use_software_profile(monkeypatch)
    posted = datetime(2026, 10, 14, tzinfo=timezone.utc)
    stub_job_batches(monkeypatch, [
        {"company": "Local", "title": "Software Engineer (Remote)", "url": "https://x/ok", "location": "Lisboa",
         "work_mode": "remote", "posted_at": posted, "industry": ["Full-time", "software engineering"],
         "excerpt": "Python and React " * 40, "source": "ITJobs",
         "schedule": "full_time", "contract": "Contrato sem termo"},
    ])
    result = analysis_from_response(client.post(
        "/career-profile/analyze",
        data={"location": "Portugal"},
        files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ))
    suggestion = result["suggestions"][0]
    assert suggestion["work_mode"] == "remote"
    assert suggestion["posted_at"] == posted.isoformat()
    assert suggestion["schedule"] == "full_time"
    assert suggestion["contract"] == "Contrato sem termo"
    assert suggestion["tags"] == ["software engineering"]
    assert len(suggestion["summary"]) <= 321
    assert suggestion["summary"].endswith("…")


def test_reviews_copy_from_teamlyzer_card_to_same_company(monkeypatch):
    use_software_profile(monkeypatch)
    stub_job_batches(monkeypatch, [
        {
            "company": "Growin",
            "title": "Software Engineer",
            "url": "https://pt.teamlyzer.com/companies/jobs/1",
            "location": "Lisboa",
            "industry": ["software engineering"],
            "excerpt": "Python and React",
            "source": "Teamlyzer",
            "rating": 2.9,
            "review_count": 19,
            "company_url": "https://pt.teamlyzer.com/companies/growin",
        },
        {
            "company": "Growin – Know to grow",
            "title": "Software Engineer",
            "url": "https://www.itjobs.pt/oferta/1/software-engineer",
            "location": "Remote",
            "work_mode": "remote",
            "industry": ["software engineering"],
            "excerpt": "Python and React",
            "source": "ITJobs",
        },
    ])
    events = list(analyze_cv_stream(
        "cv.docx",
        create_cv_docx(),
        profile={**SOFTWARE_PROFILE, "languages": ["Português"]},
    ))
    last_jobs = [event for event in events if event["type"] == "jobs"][-1]["suggestions"]
    by_source = {item["source"]: item for item in last_jobs}
    assert by_source["ITJobs"]["rating"] == 2.9
    assert by_source["ITJobs"]["review_count"] == 19
    assert by_source["ITJobs"]["company_url"] == "https://pt.teamlyzer.com/companies/growin"


def test_reviews_lookup_fills_companies_missing_from_teamlyzer(monkeypatch):
    use_software_profile(monkeypatch)
    stub_job_batches(monkeypatch, [
        {
            "company": "Growin",
            "title": "Software Engineer",
            "url": "https://www.itjobs.pt/oferta/1/software-engineer",
            "location": "Lisboa",
            "industry": ["software engineering"],
            "excerpt": "Python and React",
            "source": "ITJobs",
        },
    ])
    monkeypatch.setattr(
        "app.services.JobSearch.lookup_company",
        lambda name: {
            "rating": 2.9,
            "review_count": 130,
            "company_url": "https://pt.teamlyzer.com/companies/growin",
        } if name == "Growin" else None,
    )
    events = list(analyze_cv_stream(
        "cv.docx",
        create_cv_docx(),
        profile={**SOFTWARE_PROFILE, "languages": ["Português"]},
    ))
    last_jobs = [event for event in events if event["type"] == "jobs"][-1]["suggestions"]
    assert last_jobs[0]["rating"] == 2.9
    assert last_jobs[0]["review_count"] == 130
    assert last_jobs[0]["company_url"] == "https://pt.teamlyzer.com/companies/growin"


def test_cv_keeps_jobs_when_the_required_language_is_in_the_cv(client, monkeypatch):
    monkeypatch.setattr(
        "app.services.CareerProfile.chat_json",
        lambda *args, **kwargs: {**SOFTWARE_PROFILE, "languages": ["Inglês", "English", "Francês", "French"]},
    )
    stub_job_batches(monkeypatch, [
        {"company": "We are META", "title": ".Net Developer (French Speaker)", "url": "https://x/fr",
         "location": "Porto", "industry": ["software engineering"], "excerpt": "Python and React", "source": "Net-Empregos"},
    ])
    result = analysis_from_response(client.post(
        "/career-profile/analyze",
        data={"location": "Portugal"},
        files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ))
    assert [suggestion["company"] for suggestion in result["suggestions"]] == ["We are META"]


def test_analyze_omits_companies_already_in_outreach(client, monkeypatch):
    monkeypatch.setattr("app.controllers.CareerProfile.enrich_company", lambda *args, **kwargs: None)
    use_software_profile(monkeypatch)
    stub_job_batches(monkeypatch, [
        {
            "company": "Critical Manufacturing",
            "title": "Software Engineer",
            "url": "https://jobicy.com/jobs/example",
            "location": "Portugal",
            "industry": ["Software Engineering"],
            "excerpt": "Python and React experience",
            "source": "Jobicy",
        },
        {
            "company": "QuickNode",
            "title": "Software Engineer",
            "url": "https://jobicy.com/jobs/quicknode-portugal",
            "location": "Portugal",
            "industry": ["Software Engineering"],
            "excerpt": "Python and React experience",
            "source": "Jobicy",
        },
    ])
    client.post("/career-profile/suggestions/save", json={
        "company": "Critical Manufacturing",
        "title": "Software Engineer",
        "url": "https://jobicy.com/jobs/example",
        "source": "Jobicy",
    })
    result = analysis_from_response(client.post(
        "/career-profile/analyze",
        data={"location": "Portugal"},
        files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ))
    companies = [item["company"] for item in result["suggestions"]]
    assert "Critical Manufacturing" not in companies
    assert "QuickNode" in companies
    assert result["already_in_list_count"] == 1


def test_repeating_a_search_returns_the_same_jobs(client, monkeypatch):
    use_software_profile(monkeypatch)
    stub_job_batches(monkeypatch, [
        {"company": "Acme", "title": "Software Engineer", "url": "https://x/a",
         "location": "Lisboa", "industry": ["Software Engineering"], "excerpt": "Python", "source": "ITJobs"},
        {"company": "Beta", "title": "Software Engineer", "url": "https://x/b",
         "location": "Porto", "industry": ["Software Engineering"], "excerpt": "Python", "source": "ITJobs"},
    ])

    def companies():
        result = analysis_from_response(client.post(
            "/career-profile/analyze",
            data={"location": "Portugal"},
            files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        ))
        return sorted(item["company"] for item in result["suggestions"])

    assert companies() == ["Acme", "Beta"]
    assert companies() == ["Acme", "Beta"]


def test_analyze_hides_jobs_saved_before(client, monkeypatch):
    monkeypatch.setattr("app.controllers.CareerProfile.enrich_company", lambda *args, **kwargs: None)
    use_software_profile(monkeypatch)
    stub_job_batches(monkeypatch, [
        {"company": "Acme", "title": "Software Engineer", "url": "https://x/saved",
         "location": "Lisboa", "industry": ["Software Engineering"], "excerpt": "Python", "source": "ITJobs"},
        {"company": "Beta", "title": "Software Engineer", "url": "https://x/new",
         "location": "Porto", "industry": ["Software Engineering"], "excerpt": "Python", "source": "ITJobs"},
    ])
    saved = client.post("/career-profile/suggestions/save", json={
        "company": "Acme", "title": "Software Engineer", "url": "https://x/saved", "source": "ITJobs",
    }).json()
    assert client.delete(f"/companies/{saved['id_company']}/outreach").status_code in (200, 204)

    result = analysis_from_response(client.post(
        "/career-profile/analyze",
        data={"location": "Portugal"},
        files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ))
    assert [item["company"] for item in result["suggestions"]] == ["Beta"]


def test_draft_runs_graph_and_skips_low_fit(client, monkeypatch):
    answers = iter([
        {"fit_score": 80, "fit_reason": "Boa correspondência.", "subject": "Application", "body": "Hello, CV attached."},
        {"fit_score": 10, "fit_reason": "Pouca relação."},
    ])
    monkeypatch.setattr("app.ai.outreach_graph.chat_json", lambda *args, **kwargs: next(answers))
    payload = {"company": "Acme", "title": "Backend Engineer", "cv_skills": ["python"], "matched_skills": ["python"]}

    good = client.post("/career-profile/suggestions/draft", json=payload).json()
    assert good["skipped"] is False
    assert good["subject"] == "Application" and good["fit_score"] == 80

    weak = client.post("/career-profile/suggestions/draft", json=payload).json()
    assert weak["skipped"] is True and weak["body"] is None


def test_draft_force_ignores_low_fit(client, monkeypatch):
    monkeypatch.setattr("app.ai.outreach_graph.chat_json", lambda *args, **kwargs: {
        "fit_score": 10, "fit_reason": "Pouca relação.", "subject": "Candidatura", "body": "Olá.",
    })
    result = client.post(
        "/career-profile/suggestions/draft",
        json={"company": "Acme", "title": "Engineer", "cv_skills": ["python"], "force": True},
    ).json()
    assert result["skipped"] is False and result["subject"] == "Candidatura"


def test_draft_without_job_title_does_not_invent_software_engineer(client, monkeypatch):
    prompts = []

    def fake(system, user):
        prompts.append(user)
        return {
            "fit_score": 80,
            "fit_reason": "Perfil alinhado.",
            "subject": "Candidatura",
            "body": "Olá, CV em anexo.",
        }

    monkeypatch.setattr("app.ai.outreach_graph.chat_json", fake)
    result = client.post(
        "/career-profile/suggestions/draft",
        json={"company": "Clínica Norte", "cv_skills": ["psicologia clínica"], "force": True},
    ).json()
    assert result["skipped"] is False
    assert all("Software Engineer" not in text for text in prompts)
    assert "psicologia clínica" in prompts[0].casefold()


def test_draft_returns_503_when_model_unavailable(client, monkeypatch):
    from app.ai.llm import LLMError

    def fail(*args, **kwargs):
        raise LLMError("Ollama offline")

    monkeypatch.setattr("app.ai.outreach_graph.chat_json", fail)
    response = client.post("/career-profile/suggestions/draft", json={"company": "A", "title": "B"})
    assert response.status_code == 503


def test_cv_profile_uses_the_trade_from_the_cv(client, monkeypatch):
    monkeypatch.setattr(
        "app.services.CareerProfile.chat_json",
        lambda *args, **kwargs: {
            "candidate_name": "João Silva",
            "areas": [{
                "key": "carpintaria",
                "name": "Carpintaria",
                "search_query": "carpinteiro madeira",
                "search_tag": "carpinteiro",
                "industries": ["construção"],
                "role_terms": ["carpinteiro", "carpenter"],
                "matched_skills": ["madeira", "mobiliário"],
                "confidence": "high",
                "evidence": [{"text": "Carpinteiro de mobiliário", "section": "experiencia"}],
            }],
        },
    )

    def fake_search(tag, location=None, limit=10, include_remote=True):
        assert "carpint" in tag
        yield "ITJobs", [
            {"company": "Madeiras do Norte", "title": "Carpinteiro", "url": "https://x/wood",
             "location": "Porto", "industry": ["construção"], "excerpt": "Trabalho em madeira e mobiliário", "source": "ITJobs"},
            {"company": "Tech Co", "title": "Software Engineer", "url": "https://x/dev",
             "location": "Lisboa", "industry": ["software"], "excerpt": "Python and React", "source": "Jobicy"},
        ]

    monkeypatch.setattr("app.services.JobSearch.fetch_source_batches", fake_search)
    result = analysis_from_response(client.post(
        "/career-profile/analyze",
        data={"location": "Portugal"},
        files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ))

    assert result["areas"][0]["name"] == "Carpintaria"
    assert [item["company"] for item in result["suggestions"]] == ["Madeiras do Norte"]


def test_cv_profile_matches_clinical_psychology_jobs(client, monkeypatch):
    monkeypatch.setattr(
        "app.services.CareerProfile.chat_json",
        lambda *args, **kwargs: {
            "candidate_name": "Mauro Moreira",
            "areas": [{
                "key": "psicologia_clinica",
                "name": "Psicólogo clínico",
                "search_query": "psicólogo clínico",
                "search_tag": "clinical psychologist",
                "industries": ["saúde"],
                "role_terms": ["psicólogo clínico", "psicólogo", "clinical psychologist"],
                "matched_skills": ["avaliação psicológica", "supervisão"],
                "confidence": "high",
                "evidence": [{"text": "Psicólogo clínico", "section": "experiencia"}],
            }],
        },
    )

    def fake_search(tag, location=None, limit=10, include_remote=True):
        assert "psicologo" in tag
        assert "clinical" not in tag
        yield "Net-Empregos", [
            {"company": "Clínica Norte", "title": "Psicólogo Clínico - LISBOA", "url": "https://www.net-empregos.com/1/",
             "location": "Lisboa", "industry": ["Saúde"], "excerpt": "Saúde / Medicina", "source": "Net-Empregos"},
            {"company": "Tech Co", "title": "Software Engineer", "url": "https://x/dev",
             "location": "Lisboa", "industry": ["software"], "excerpt": "Python", "source": "ITJobs"},
        ]

    monkeypatch.setattr("app.services.JobSearch.fetch_source_batches", fake_search)
    result = analysis_from_response(client.post(
        "/career-profile/analyze",
        data={"location": "Portugal"},
        files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ))
    assert result["areas"][0]["name"] == "Psicólogo clínico"
    assert [item["company"] for item in result["suggestions"]] == ["Clínica Norte"]


def test_cv_profile_falls_back_to_cv_text_when_ollama_is_down(client, monkeypatch):
    monkeypatch.setattr(
        "app.services.CareerProfile.chat_json",
        lambda *args, **kwargs: (_ for _ in ()).throw(LLMError("offline")),
    )
    stub_job_batches(monkeypatch, [])
    result = analysis_from_response(client.post(
        "/career-profile/analyze",
        data={"location": "Portugal"},
        files={"file": ("cv.docx", create_cv_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ))
    assert result["areas"]
    assert result["llm_used"] is False
    assert "software" in result["areas"][0]["name"].casefold() or "python" in " ".join(result["areas"][0]["matched_skills"]).casefold()


def test_portugal_search_keeps_lisboa_and_porto():
    from app.services.JobSearch import _location_priority

    assert _location_priority("Lisboa", "Portugal") == 0
    assert _location_priority("Lisboa, Porto", "Portugal") == 0
    assert _location_priority("Porto", "Portugal") == 0
    assert _location_priority("Leiria", "Portugal") == 0
    assert _location_priority("Oeiras", "Portugal") == 0
    assert _location_priority("Guimarães", "Portugal") == 0
    assert _location_priority("Sintra", "Portugal") == 0
    assert _location_priority("Madrid", "Portugal") is None
    assert _location_priority("London", "Portugal") is None
    assert _location_priority("Madrid", "Espanha") == 0
    assert _location_priority("Remote, Worldwide", "Portugal") == 3
    assert _location_priority("Europe", "Portugal") == 2
    assert _location_priority("Berlin, Germany", "Alemanha") == 0
    assert _location_priority("München", "Alemanha") == 0
    assert _location_priority("Berlin, Germany", "Portugal") is None
    assert _location_priority("Lisboa", "Alemanha") is None
    assert _location_priority("Braga (Regime: Híbrido)", "Braga", "hybrid") == 0
    assert _location_priority("Braga", "Braga") == 0
    assert _location_priority("Braga, Portugal", "Braga") == 0
    assert _location_priority("Lisboa", "Braga") is None
    assert _location_priority("Lisboa, Portugal", "Braga") is None
    assert _location_priority("Porto", "Braga") is None
    assert _location_priority("Portugal", "Braga") == 2
    assert _location_priority("Remote", "Braga") == 3
    assert _location_priority("Remote, Worldwide", "Braga") == 3
    assert _location_priority("Europe", "Braga") == 2
    assert _location_priority("Remote, USA", "Braga") is None
    assert _location_priority("Lisboa", "Braga", "remote") == 3


def test_extract_candidate_name():
    from app.services.CareerProfile import extract_candidate_name

    assert extract_candidate_name("Miguel Santos\nSoftware Engineer") == "Miguel Santos"
    assert extract_candidate_name("\n  Ana Maria Costa  \nPorto") == "Ana Maria Costa"
    assert extract_candidate_name("Curriculum Vitae\nMiguel Santos") == "Curriculum Vitae"
    assert extract_candidate_name("miguel@mail.com\nMiguel") is None
    assert extract_candidate_name("Experiência Profissional 2020-2024") is None


def test_extract_profile_rejects_name_as_area_and_hobby_skills(monkeypatch):
    monkeypatch.setattr(
        "app.services.CareerProfile.chat_json",
        lambda *args, **kwargs: {
            "candidate_name": "Miguel Silva",
            "areas": [{
                "name": "Miguel Silva",
                "search_tag": "Miguel Silva",
                "role_terms": ["Miguel Silva", "Software Developer", "programador"],
                "matched_skills": (
                    "React SQL PostgreSQL FastAPI Pydantic LANGUAGES Portuguese English "
                    "HOBBIES Music Photography"
                ),
                "evidence": [
                    {"text": "Miguel Silva", "section": "perfil"},
                    {"text": "React", "section": "competencias"},
                    {"text": "Mobile Application Development, Polytechnic Institute of Braga", "section": "formacao"},
                ],
            }],
        },
    )
    from app.services.CareerProfile import extract_profile

    name, areas, llm_used, languages = extract_profile("Miguel Silva\nSoftware Developer\nReact, SQL\n")
    assert llm_used is True
    skills = [skill.casefold() for skill in areas[0]["matched_skills"]]

    assert name == "Miguel Silva"
    assert areas[0]["name"] == "Software Developer"
    assert "react" in skills and "sql" in skills
    assert "music" not in skills and "portuguese" not in skills
    assert "miguel" not in areas[0]["search_tag"].casefold()
    assert languages == []


def test_cv_profile_fallback_skips_candidate_name(client, monkeypatch):
    monkeypatch.setattr(
        "app.services.CareerProfile.chat_json",
        lambda *args, **kwargs: (_ for _ in ()).throw(LLMError("offline")),
    )
    stub_job_batches(monkeypatch, [])
    document = Document()
    document.add_paragraph("Miguel Silva")
    document.add_paragraph("Software Developer")
    document.add_paragraph("React, SQL, PostgreSQL, FastAPI")
    document.add_paragraph("HOBBIES")
    document.add_paragraph("Music Photography")
    document.add_paragraph("Formação Académica")
    document.add_paragraph("Mobile Application Development")
    content = BytesIO()
    document.save(content)

    result = analysis_from_response(client.post(
        "/career-profile/analyze",
        data={"location": "Portugal"},
        files={"file": ("cv.docx", content.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ))

    assert result["candidate_name"] == "Miguel Silva"
    assert result["areas"][0]["name"] != "Miguel Silva"
    assert "software" in result["areas"][0]["name"].casefold()
    skills = " ".join(result["areas"][0]["matched_skills"]).casefold()
    assert "react" in skills
    assert "music" not in skills


def test_draft_signs_with_candidate_name(client, monkeypatch):
    monkeypatch.setattr("app.ai.outreach_graph.chat_json", lambda *args, **kwargs: {
        "fit_score": 90, "fit_reason": "ok", "subject": "Candidatura", "body": "Olá,\nCV em anexo.\n[Nome]",
    })
    result = client.post(
        "/career-profile/suggestions/draft",
        json={"company": "Acme", "title": "Dev", "language": "pt", "candidate_name": "Miguel Santos"},
    ).json()
    assert result["body"].endswith("Miguel Santos")


def test_places_autocomplete(client, monkeypatch):
    monkeypatch.setattr(
        "app.controllers.CareerProfile.search_places",
        lambda q: [{"label": "Braga, Portugal", "name": "Braga", "country": "Portugal", "kind": "city"}] if q == "Braga" else [],
    )
    listed = client.get("/career-profile/places?q=Braga")
    assert listed.status_code == 200
    assert listed.json()[0]["label"] == "Braga, Portugal"
    assert client.get("/career-profile/places?q=x").json() == []
    assert client.get("/career-profile/places").json() == []


def test_places_keeps_names_that_match_the_query(monkeypatch):
    monkeypatch.setattr("app.services.Places.get_cached_json", lambda *a, **k: [
        {"addresstype": "city", "name": "Braga", "address": {"city": "Braga", "country": "Portugal"}},
        {"addresstype": "city", "name": "Caprino Veronese", "address": {"city": "Caprino Veronese", "country": "Itália"}},
        {"addresstype": "country", "name": "Portugal", "address": {"country": "Portugal"}},
    ])
    from app.services.Places import search_places
    labels = [place["label"] for place in search_places("Braga")]
    assert labels == ["Braga, Portugal"]
