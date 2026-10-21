from io import BytesIO
from xml.etree import ElementTree as ET
import html
import logging
import re
import unicodedata
import zipfile

from docx import Document
from pypdf import PdfReader

from app.ai.llm import LLMError, chat_json

logger = logging.getLogger(__name__)


MAX_CV_BYTES = 5 * 1024 * 1024
MAX_SEARCH_AREAS = 3

IMAGE_KINDS = {"jpg", "jpeg", "png", "webp"}


SECTION_NAMES = {
    "experiencia": ("experiência", "experiencia profissional", "professional experience", "experience"),
    "formacao": ("formação", "formacao academica", "educação", "education", "academic background"),
    "competencias": ("competências", "competencias", "skills", "technical skills"),
    "perfil": ("perfil", "profile", "summary", "resumo", "sobre mim", "about me"),
}


def _normalize(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    normalized = "".join(char for char in normalized if not unicodedata.combining(char))
    normalized = normalized.casefold().replace(".", " ")
    return re.sub(r"\s+", " ", normalized).strip()


def _has_term(normalized_text: str, term: str) -> bool:
    folded = _normalize(term)
    if not folded:
        return False
    return re.search(rf"(?<![a-z0-9]){re.escape(folded)}(?![a-z0-9])", normalized_text) is not None

STOPWORDS = {
    "a", "as", "o", "os", "um", "uma", "de", "da", "do", "das", "dos", "e", "em",
    "com", "para", "por", "na", "no", "the", "and", "or", "to", "of", "in", "on",
    "at", "for", "with", "cv", "curriculum", "vitae", "experiencia", "formacao",
}

PROFILE_SYSTEM = (
    "Extract a job-search profile from the CV. Reply only with JSON. "
    "candidate_name is the person's name. areas[].name is the profession or role "
    "(e.g. Software Developer), never the person's name. "
    "matched_skills is 5 to 12 short tools or skills (React, Python, SQL), never "
    "hobbies or full sentences. "
    "languages lists spoken languages written in the CV, as written and in English "
    "(Inglês and English). Empty if the CV does not list any. Never infer a language. "
    "search_tag is 2 to 4 words in the CV language to search job boards "
    "(psicólogo clínico, carpinteiro, software developer). "
    "role_terms are usual job titles for this trade in the CV language and in English "
    "(programador and software developer, or carpinteiro and carpenter). "
    "Do not assume a technology career: a carpenter must be labelled carpenter. "
    "Use only facts written in the CV. "
    'Schema: {"candidate_name": string|null, "languages": [string], '
    '"areas":[{"key": slug, "name": string, '
    '"search_query": string, "search_tag": string, "industries": [string], '
    '"role_terms": [string], "matched_skills": [string], '
    '"confidence": "high"|"medium"|"low", '
    '"evidence": [{"text": quote, "section": "experiencia"|"formacao"|"competencias"|"perfil"}]}]}. '
    "Return 1 to 3 areas."
)

SKILL_SECTION_STOP = {"hobbies", "hobby", "languages", "idiomas", "interests", "interesses"}


def _slug(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", _normalize(value)).strip("_")
    return slug[:60] or "perfil"


def _looks_like_person_name(line: str) -> bool:
    words = line.split()
    return (
        2 <= len(words) <= 5
        and len(line) <= 60
        and not re.search(r"[\d@:/|,]", line)
        and all(word[0].isalpha() for word in words)
    )


def _same_person(line: str, candidate_name: str | None) -> bool:
    return bool(candidate_name) and _normalize(line) == _normalize(candidate_name)


def _is_dump(line: str) -> bool:
    return len(line) > 80 or len(line.split()) > 8


def _usable_title(line: str, candidate_name: str | None) -> bool:
    text = line.strip()
    return bool(text) and not _same_person(text, candidate_name) and not _is_dump(text)


def _string_list(value) -> list[str]:
    if isinstance(value, str):
        value = re.split(r"[,;/|•·]", value)
    if not isinstance(value, list):
        return []
    return list(dict.fromkeys(item.strip() for item in value if isinstance(item, str) and item.strip()))


def _skill_pieces(values) -> list[str]:
    pieces = []
    for skill in _string_list(values):
        if len(skill) > 40:
            pieces.extend(part.strip() for part in re.split(r"[\s,/]+", skill) if part.strip())
        else:
            pieces.append(skill)
    return pieces


def _clean_skills(values, candidate_name: str | None = None) -> list[str]:
    name_parts = {_normalize(part) for part in _tokens(candidate_name or "")}
    cleaned = []
    for skill in _skill_pieces(values):
        folded = _normalize(skill)
        if folded in SKILL_SECTION_STOP or folded.startswith("hobby"):
            break
        if folded in STOPWORDS or folded in name_parts:
            continue
        if not (2 <= len(skill) <= 40):
            continue
        cleaned.append(skill)
        if len(cleaned) >= 12:
            break
    return cleaned


def _finalize_area(area: dict, candidate_name: str | None) -> dict:
    area["matched_skills"] = _clean_skills(area["matched_skills"], candidate_name)
    if not _usable_title(area["name"], candidate_name):
        replacement = next((term for term in area["role_terms"] if _usable_title(term, candidate_name)), None)
        if not replacement:
            replacement = next(
                (item["text"] for item in area["evidence"] if _usable_title(item["text"], candidate_name)),
                None,
            )
        if replacement:
            area["name"] = replacement[:80]
            area["key"] = _slug(replacement)
        elif area["matched_skills"]:
            area["name"] = " ".join(area["matched_skills"][:3])[:80]
            area["key"] = _slug(area["name"])
    tag = area.get("search_tag") or ""
    if not _usable_title(tag, candidate_name):
        source = next((term for term in area["role_terms"] if _usable_title(term, candidate_name)), area["name"])
        area["search_tag"] = " ".join(_tokens(source)[:4]) or " ".join(area["matched_skills"][:4])
    return area


def _evidence_items(value) -> list[dict]:
    items = []
    if not isinstance(value, list):
        return items
    for item in value[:4]:
        if not isinstance(item, dict):
            continue
        text = str(item.get("text") or "").strip()
        section = str(item.get("section") or "perfil").strip()
        if section not in SECTION_NAMES:
            section = "perfil"
        if text:
            items.append({"text": text[:240], "section": section})
    return items


def _area_from_dict(item: dict) -> dict | None:
    name = str(item.get("name") or "").strip()
    if not name:
        return None
    skills = _string_list(item.get("matched_skills") or item.get("skills"))
    roles = _string_list(item.get("role_terms") or item.get("roles")) or [name]
    search_query = str(item.get("search_query") or "").strip() or " ".join([name, *skills[:3]])
    search_tag = str(item.get("search_tag") or "").strip() or name
    confidence = item.get("confidence") if item.get("confidence") in {"high", "medium", "low"} else "medium"
    return {
        "key": _slug(str(item.get("key") or name)),
        "name": name[:80],
        "search_query": search_query[:180],
        "search_tag": search_tag[:80],
        "industries": _string_list(item.get("industries")),
        "role_terms": roles,
        "matched_skills": skills,
        "confidence": confidence,
        "evidence": _evidence_items(item.get("evidence")),
    }


def _areas_from_payload(payload: dict) -> list[dict]:
    raw = payload.get("areas") if isinstance(payload, dict) else None
    if not isinstance(raw, list):
        return []
    areas = []
    for item in raw:
        if not isinstance(item, dict):
            continue
        area = _area_from_dict(item)
        if area:
            areas.append(area)
        if len(areas) >= MAX_SEARCH_AREAS:
            break
    return areas


def _cv_sections(text: str) -> list[tuple[str, str]]:
    contextual_lines: list[tuple[str, str]] = []
    current_section = "perfil"
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        normalized_line = _normalize(line)
        for section, headings in SECTION_NAMES.items():
            if normalized_line in {_normalize(heading) for heading in headings}:
                current_section = section
                break
        else:
            contextual_lines.append((line, current_section))
    return contextual_lines


def _tokens(value: str) -> list[str]:
    return [
        token
        for token in re.findall(r"[A-Za-zÀ-ÿ0-9+#.]{3,}", value)
        if _normalize(token) not in STOPWORDS
    ]


def _areas_from_cv_text(text: str, candidate_name: str | None = None) -> list[dict]:
    """Fallback if Ollama is down: still uses the CV text, not a fixed career list."""
    sections = _cv_sections(text)
    if not sections:
        return []

    experience = [
        line for line, section in sections
        if section in {"experiencia", "perfil"} and not _same_person(line, candidate_name)
    ]
    formation = [line for line, section in sections if section == "formacao"]
    skill_lines = [line for line, section in sections if section == "competencias"]
    role_line = next((line for line in experience if not _is_dump(line)), None)
    if not role_line:
        role_line = formation[0] if formation else (experience[0] if experience else sections[0][0])
    skill_sources = skill_lines or [line for line in experience if "," in line or _is_dump(line)]
    raw_skills = skill_sources or experience[:2]
    skills = _skill_pieces(raw_skills) or _tokens(" ".join(raw_skills))
    roles = [role_line, *formation[:1], *_tokens(role_line)[:6]]
    evidence = []
    for line, section in sections:
        if line in {role_line, *formation[:1], *skill_lines[:1]}:
            evidence.append({"text": line[:240], "section": section})
        if len(evidence) >= 4:
            break
    return [_area_from_dict({
        "name": role_line[:80],
        "search_query": " ".join([role_line, *skills[:3]]),
        "search_tag": " ".join(_tokens(role_line)[:4]) or role_line,
        "role_terms": roles,
        "matched_skills": skills,
        "confidence": "medium",
        "evidence": evidence or [{"text": role_line[:240], "section": "perfil"}],
    })]


def extract_profile(text: str) -> tuple[str | None, list[dict], bool, list[str]]:
    try:
        payload = chat_json(PROFILE_SYSTEM, text[:7_000])
        areas = _areas_from_payload(payload)
        if areas:
            name = str(payload.get("candidate_name") or "").strip() or extract_candidate_name(text)
            return name, [_finalize_area(area, name) for area in areas], True, _string_list(payload.get("languages"))
    except LLMError as error:
        logger.warning("CV profile LLM failed: %s", error.__cause__ or error)
    name = extract_candidate_name(text)
    return name, [_finalize_area(area, name) for area in _areas_from_cv_text(text, name)], False, []

def _profile_from_stored(profile: dict | None) -> tuple[str | None, list[dict], bool, list[str]] | None:
    if not profile:
        return None
    name = str(profile.get("candidate_name") or "").strip() or None
    areas = [_finalize_area(area, name) for area in _areas_from_payload(profile)]
    if not areas:
        return None
    reliable = bool(profile.get("llm_used") or profile.get("edited"))
    return name, areas, reliable, _string_list(profile.get("languages"))


def edit_profile(profile: dict | None, candidate_name: str | None, edits: list[dict]) -> dict:
    """Aplica as correções do utilizador. Uma área com outro nome deixa de herdar os termos que a IA lhe deu."""
    current = {area["name"].casefold(): area for area in _areas_from_payload(profile or {})}
    areas = []
    for edit in edits:
        base = current.get(edit["name"].casefold())
        area = {**base, "matched_skills": edit["skills"]} if base else _area_from_dict(
            {"name": edit["name"], "matched_skills": edit["skills"]}
        )
        areas.append(_finalize_area(area, candidate_name))
    return {
        "candidate_name": candidate_name,
        "areas": areas,
        "languages": _string_list((profile or {}).get("languages")),
        "llm_used": bool((profile or {}).get("llm_used")),
        "edited": True,
    }

_RTF_TOKEN = re.compile(
    r"\\([a-z]{1,32})(-?\d{1,10})?[ ]?|\\'([0-9a-f]{2})|\\([^a-z])|([{}])|[\r\n]+|(.)",
    re.I,
)


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _zip_kind(content: bytes) -> str | None:
    try:
        names = zipfile.ZipFile(BytesIO(content)).namelist()
    except zipfile.BadZipFile:
        return None
    if "word/document.xml" in names:
        return "docx"
    if "content.xml" in names:
        return "odt"
    return None


_ocr = None


def _image_kind(content: bytes) -> str | None:
    if content.startswith(b"\xff\xd8\xff"):
        return "jpg"
    if content.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if content[:4] == b"RIFF" and content[8:12] == b"WEBP":
        return "webp"
    return None


def _ocr_engine():
    global _ocr
    if _ocr is None:
        from rapidocr import RapidOCR

        _ocr = RapidOCR(params={"Rec.lang_type": "pt", "Global.log_level": "error"})
    return _ocr


def _image_text(content: bytes) -> str:
    try:
        result = _ocr_engine()(content)
    except Exception as error:
        logger.warning("CV OCR failed: %s", error)
        raise ValueError("Não foi possível ler o CV em imagem. Tenta outra foto ou um PDF.") from error
    texts = getattr(result, "txts", None) or ()
    return "\n".join(part.strip() for part in texts if isinstance(part, str) and part.strip())


def _format_of(filename: str, content: bytes) -> str:
    extension = filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    if content.startswith(b"%PDF"):
        return "pdf"
    if content.lstrip().startswith(b"{\\rtf"):
        return "rtf"
    if content.startswith(b"PK"):
        kind = _zip_kind(content)
        if kind:
            return kind
    if content.startswith(b"\xd0\xcf\x11\xe0") or extension == "doc":
        raise ValueError("Ficheiros .doc antigos não são suportados. Grava o CV em PDF, DOCX ou ODT.")
    kind = _image_kind(content)
    if kind:
        return kind
    if extension in IMAGE_KINDS:
        return "jpg" if extension == "jpeg" else extension
    if extension in {"pdf", "docx", "odt", "rtf", "txt", "md", "html", "htm"}:
        return extension
    raise ValueError("Formato não suportado. Usa PDF, DOCX, ODT, RTF, JPG, PNG ou um ficheiro de texto.")


def _docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + table_cells)


def _odt_text(content: bytes) -> str:
    with zipfile.ZipFile(BytesIO(content)) as archive:
        xml = archive.read("content.xml")
    parts = []
    for node in ET.fromstring(xml).iter():
        if node.text and node.text.strip():
            parts.append(node.text.strip())
        if node.tail and node.tail.strip():
            parts.append(node.tail.strip())
    return "\n".join(parts)


def _rtf_text(content: bytes) -> str:
    stack: list[bool] = []
    ignorable = False
    out: list[str] = []
    for control, _number, hexcode, escaped, brace, plain in _RTF_TOKEN.findall(_decode_text(content)):
        if brace == "{":
            stack.append(ignorable)
        elif brace == "}":
            ignorable = stack.pop() if stack else False
        elif hexcode and not ignorable:
            out.append(bytes.fromhex(hexcode).decode("cp1252", errors="ignore"))
        elif control in {"*", "fonttbl", "colortbl", "stylesheet", "info", "pict"}:
            ignorable = True
        elif control in {"par", "line"} and not ignorable:
            out.append("\n")
        elif control == "tab" and not ignorable:
            out.append("\t")
        elif escaped and not ignorable:
            out.append(escaped)
        elif plain and not ignorable:
            out.append(plain)
    return "".join(out)


def _html_text(content: bytes) -> str:
    raw = _decode_text(content)
    raw = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", raw)
    raw = re.sub(r"<[^>]+>", "\n", raw)
    return re.sub(r"\n{3,}", "\n\n", html.unescape(raw))


def extract_cv_text(filename: str, content: bytes) -> str:
    kind = _format_of(filename, content)
    if kind == "pdf":
        pages = PdfReader(BytesIO(content)).pages
        text = "\n".join(page.extract_text() or "" for page in pages)
    elif kind == "docx":
        text = _docx_text(content)
    elif kind == "odt":
        text = _odt_text(content)
    elif kind == "rtf":
        text = _rtf_text(content)
    elif kind in {"html", "htm"}:
        text = _html_text(content)
    elif kind in {"jpg", "png", "webp"}:
        text = _image_text(content)
    else:
        text = _decode_text(content)

    text = text.strip()
    if not text:
        if kind in {"jpg", "png", "webp"}:
            raise ValueError("Não foi possível ler texto na imagem. Tira uma foto mais nítida ou usa PDF/DOCX.")
        raise ValueError("Não foi possível extrair texto. Ficheiros digitalizados como imagem não são suportados.")

    return text[:100_000]


def extract_candidate_name(text: str) -> str | None:
    """O nome costuma ser a primeira linha do CV; rejeita linhas que não pareçam um nome."""
    for line in text.splitlines():
        line = line.strip()
        if line:
            return line if _looks_like_person_name(line) else None
    return None
